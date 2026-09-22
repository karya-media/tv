"""Overrides spreadsheet writer (.xlsx) using openpyxl.

Exports every channel's current metadata plus two blank "correction"
columns (New tvg-id / New group-title) a human fills in by hand. The
same file, once edited, is read back by overrides_excel_reader and
applied via ApplyOverridesUseCase - a manual-correction workflow for
metadata mistakes that are easier to spot and fix by eye in a
spreadsheet than by editing category .m3u files or Python code
directly (see the TV9-tagged-as-India case fixed earlier in this
project's history).

URL is the join key back to a channel (same key
MergePlaylistsUseCase already uses for duplicate detection) - included
as its own column so this file is self-contained and inspectable, but
it isn't meant to be edited.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from iptv_manager.domain.entities.playlist import Playlist

_HEADER_FONT = Font(bold=True)
_HEADERS = [
    "Source File",
    "Channel Name",
    "URL",
    "tvg-id",
    "New tvg-id",
    "group-title",
    "New group-title",
]


class OverridesExcelWriter:
    def write(self, playlist: Playlist, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Overrides"

        sheet.append(_HEADERS)
        for cell in sheet[1]:
            cell.font = _HEADER_FONT

        for channel in playlist:
            sheet.append(
                [
                    channel.source_category or "",
                    channel.name,
                    channel.url.raw,
                    str(channel.tvg_id),
                    "",  # New tvg-id - left blank for the user to fill in
                    str(channel.group_title),
                    "",  # New group-title - left blank for the user to fill in
                ]
            )

        self._autosize_columns(sheet)
        workbook.save(path)

    def _autosize_columns(self, sheet: Worksheet) -> None:
        for index, column_cells in enumerate(sheet.columns, start=1):
            length = max((len(str(cell.value)) for cell in column_cells if cell.value), default=0)
            column_letter = get_column_letter(index)
            sheet.column_dimensions[column_letter].width = min(max(length + 2, 10), 60)
