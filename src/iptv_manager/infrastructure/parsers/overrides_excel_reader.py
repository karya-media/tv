"""Overrides spreadsheet reader (.xlsx) using openpyxl.

Reads back the file overrides_excel_writer.OverridesExcelWriter
produces (once a human has filled in its "New tvg-id"/"New
group-title" columns), matching this project's column order/naming
exactly. See domain.entities.channel_override.ChannelOverride and
application.use_cases.apply_overrides.ApplyOverridesUseCase for how
the parsed result gets used.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from iptv_manager.domain.entities.channel_override import ChannelOverride

# Must match overrides_excel_writer._HEADERS exactly (column order).
_URL_COLUMN = 3
_NEW_TVG_ID_COLUMN = 5
_NEW_GROUP_TITLE_COLUMN = 7


class OverridesExcelReadError(ValueError):
    """Raised when the file isn't a readable/well-formed overrides
    spreadsheet."""


def read_overrides_excel(path: Path) -> list[ChannelOverride]:
    """Returns an empty list if the file doesn't exist - overrides are
    entirely opt-in."""
    if not path.exists():
        return []

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises several distinct types
        raise OverridesExcelReadError(f"could not read {path} as an .xlsx file: {exc}") from exc

    try:
        sheet = workbook["Overrides"] if "Overrides" in workbook.sheetnames else workbook.active
        if sheet is None:
            return []

        overrides: list[ChannelOverride] = []
        for row in sheet.iter_rows(min_row=2):  # skip header row
            url = _cell_text(row, _URL_COLUMN)
            if not url:
                continue
            overrides.append(
                ChannelOverride(
                    url=url,
                    new_tvg_id=_cell_text(row, _NEW_TVG_ID_COLUMN),
                    new_group_title=_cell_text(row, _NEW_GROUP_TITLE_COLUMN),
                )
            )
        return overrides
    finally:
        workbook.close()


def _cell_text(row: Any, column: int) -> str | None:
    if column > len(row):
        return None
    value = row[column - 1].value
    if value is None:
        return None
    text = str(value).strip()
    return text or None
