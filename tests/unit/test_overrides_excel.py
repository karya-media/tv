"""Unit tests for infrastructure.reports.overrides_excel_writer and
infrastructure.parsers.overrides_excel_reader - covered together since
the reader is defined to be the writer's exact counterpart (same
column order/names)."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from iptv_manager.domain.entities.channel import Channel
from iptv_manager.domain.entities.playlist import Playlist
from iptv_manager.domain.value_objects.group_title import GroupTitle
from iptv_manager.domain.value_objects.stream_url import StreamUrl
from iptv_manager.domain.value_objects.tvg_id import TvgId
from iptv_manager.infrastructure.parsers.overrides_excel_reader import (
    OverridesExcelReadError,
    read_overrides_excel,
)
from iptv_manager.infrastructure.reports.overrides_excel_writer import OverridesExcelWriter


def _channel(
    name: str, url: str, tvg_id: str | None = None, group_title: str | None = None
) -> Channel:
    return Channel(
        name=name,
        url=StreamUrl.parse(url),
        tvg_id=TvgId.parse(tvg_id),
        group_title=GroupTitle.parse(group_title),
        source_category="02_lokal",
    )


class TestOverridesExcelWriter:
    def test_writes_one_row_per_channel_plus_header(self, tmp_path: Path):
        playlist = Playlist(
            name="test",
            channels=[
                _channel("TV9", "http://x.com/1.m3u8", tvg_id="TV9.in", group_title="India;News"),
                _channel("RCTI", "http://x.com/2.m3u8", tvg_id="RCTI.id"),
            ],
        )
        output_path = tmp_path / "overrides.xlsx"
        OverridesExcelWriter().write(playlist, output_path)

        workbook = load_workbook(output_path)
        sheet = workbook["Overrides"]
        assert sheet.max_row == 3  # header + 2 channels

    def test_header_row_matches_expected_columns(self, tmp_path: Path):
        playlist = Playlist(name="test", channels=[_channel("X", "http://x.com/1.m3u8")])
        output_path = tmp_path / "overrides.xlsx"
        OverridesExcelWriter().write(playlist, output_path)

        workbook = load_workbook(output_path)
        sheet = workbook["Overrides"]
        header = [cell.value for cell in sheet[1]]
        assert header == [
            "Source File",
            "Channel Name",
            "URL",
            "tvg-id",
            "New tvg-id",
            "group-title",
            "New group-title",
        ]

    def test_new_tvg_id_and_group_title_columns_start_blank(self, tmp_path: Path):
        playlist = Playlist(
            name="test",
            channels=[_channel("TV9", "http://x.com/1.m3u8", tvg_id="TV9.in")],
        )
        output_path = tmp_path / "overrides.xlsx"
        OverridesExcelWriter().write(playlist, output_path)

        workbook = load_workbook(output_path)
        sheet = workbook["Overrides"]
        row = list(sheet.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        assert row[4] is None  # New tvg-id
        assert row[6] is None  # New group-title

    def test_source_category_is_written(self, tmp_path: Path):
        playlist = Playlist(name="test", channels=[_channel("TV9", "http://x.com/1.m3u8")])
        output_path = tmp_path / "overrides.xlsx"
        OverridesExcelWriter().write(playlist, output_path)

        workbook = load_workbook(output_path)
        sheet = workbook["Overrides"]
        row = list(sheet.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        assert row[0] == "02_lokal"


class TestReadOverridesExcel:
    def test_missing_file_returns_empty_list(self, tmp_path: Path):
        assert read_overrides_excel(tmp_path / "does-not-exist.xlsx") == []

    def test_reads_back_a_filled_in_correction(self, tmp_path: Path):
        playlist = Playlist(
            name="test",
            channels=[_channel("TV9", "http://x.com/1.m3u8", tvg_id="TV9.in")],
        )
        output_path = tmp_path / "overrides.xlsx"
        OverridesExcelWriter().write(playlist, output_path)

        # Simulate a human filling in the correction columns.
        workbook = load_workbook(output_path)
        sheet = workbook["Overrides"]
        sheet.cell(row=2, column=5, value="TV9.id")  # New tvg-id
        sheet.cell(row=2, column=7, value="Indonesia;Lokal")  # New group-title
        workbook.save(output_path)

        overrides = read_overrides_excel(output_path)
        assert len(overrides) == 1
        assert overrides[0].url == "http://x.com/1.m3u8"
        assert overrides[0].new_tvg_id == "TV9.id"
        assert overrides[0].new_group_title == "Indonesia;Lokal"

    def test_row_with_no_url_is_skipped(self, tmp_path: Path):
        playlist = Playlist(name="test", channels=[])
        output_path = tmp_path / "overrides.xlsx"
        OverridesExcelWriter().write(playlist, output_path)

        workbook = load_workbook(output_path)
        sheet = workbook["Overrides"]
        sheet.append(["", "", "", "", "SomeValue", "", ""])  # blank URL
        workbook.save(output_path)

        assert read_overrides_excel(output_path) == []

    def test_unedited_export_round_trips_to_all_empty_overrides(self, tmp_path: Path):
        playlist = Playlist(
            name="test",
            channels=[_channel("TV9", "http://x.com/1.m3u8", tvg_id="TV9.in")],
        )
        output_path = tmp_path / "overrides.xlsx"
        OverridesExcelWriter().write(playlist, output_path)

        overrides = read_overrides_excel(output_path)
        assert len(overrides) == 1
        assert overrides[0].is_empty

    def test_non_xlsx_file_raises_a_clear_error(self, tmp_path: Path):
        bad_path = tmp_path / "not_an_excel_file.xlsx"
        bad_path.write_text("this is not a real xlsx file", encoding="utf-8")
        with pytest.raises(OverridesExcelReadError):
            read_overrides_excel(bad_path)
